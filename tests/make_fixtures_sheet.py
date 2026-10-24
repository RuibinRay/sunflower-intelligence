#!/usr/bin/env python3
"""Reference values for tests/sheet.test.mjs: every formula below is written into a workbook with openpyxl,
recalculated by LibreOffice Calc (headless), and the results are saved to tests/fixtures-sheet.json.

Run:  python3 tests/make_fixtures_sheet.py [path/to/recalc.py]
Needs LibreOffice (soffice) and openpyxl. The recalculation uses a small LibreOffice Basic macro (calculateAll, then
store), the same approach as the xlsx skill's recalc.py; pass that script's path to use it instead.

Functions new in Excel 365 (XLOOKUP, XMATCH, FILTER, SORT, SORTBY, UNIQUE, SEQUENCE, TEXTBEFORE, TEXTAFTER) are not
in LibreOffice 24.2, so tests/sheet.test.mjs checks them against Microsoft's documented examples instead. Where
LibreOffice and Excel are documented to differ, the case is listed in EXCEL_DIFFERS and the test uses Excel's value.
"""
import datetime as dt, json, os, shutil, subprocess, sys, tempfile
from pathlib import Path
from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
serial = lambda y, m, d: (dt.date(y, m, d) - dt.date(1899, 12, 30)).days

ORDERS = [  # date, channel, product, qty, price
    ((2026, 3, 2), 'Online', 'House Blend', 2, 14),
    ((2026, 3, 3), 'Café', 'Flat White', 3, 4.2),
    ((2026, 3, 5), 'Online', 'Decaf Blend', 1, 15),
    ((2026, 3, 9), 'Wholesale', 'House Blend', 20, 9.5),
    ((2026, 3, 12), 'Café', 'Pastry', 4, 3.25),
    ((2026, 3, 14), 'Online', 'House Blend', 3, 14),
    ((2026, 3, 15), 'Café', 'Flat White', 2, 4.2),
    ((2026, 3, 18), 'Wholesale', 'Single Origin', 10, 12),
    ((2026, 3, 21), 'Online', 'Single Origin', 1, 18),
    ((2026, 3, 25), 'Café', 'Pastry', 6, 3.25),
    ((2026, 3, 28), 'Online', 'Decaf Blend', 2, 15),
    ((2026, 3, 31), 'Wholesale', 'House Blend', 15, 9.5),
]
CELLS = {'A1': 'Date', 'B1': 'Channel', 'C1': 'Product', 'D1': 'Qty', 'E1': 'Price', 'F1': 'Revenue'}
for i, (d, ch, pr, q, p) in enumerate(ORDERS, start=2):
    CELLS.update({f'A{i}': serial(*d), f'B{i}': ch, f'C{i}': pr, f'D{i}': q, f'E{i}': p, f'F{i}': round(q * p, 2)})
# mixed types for coercion rules (H4 is left empty)
CELLS.update({'H1': 10, 'H2': '5', 'H3': True, 'H5': 'abc', 'H6': -3.5, 'H7': '  Lantern   coffee ', 'H8': 0})
# an approximate-match rate table, a quarterly row, a descending list, a list with tied modes
CELLS.update({'J1': 'Spend', 'K1': 'Discount', 'J2': 0, 'K2': 0, 'J3': 100, 'K3': 0.02, 'J4': 250, 'K4': 0.05, 'J5': 500, 'K5': 0.08})
CELLS.update({'M1': 'Q1', 'N1': 'Q2', 'O1': 'Q3', 'P1': 'Q4', 'M2': 120, 'N2': 135, 'O2': 150, 'P2': 160})
CELLS.update({'R1': 100, 'R2': 75, 'R3': 50, 'R4': 25})
CELLS.update({f'S{i}': v for i, v in enumerate([3, 5, 5, 7, 3, 9, 2, 1], start=1)})
# B04: two samples (bag weights), a contingency table with its expected counts, a one-row goodness-of-fit pair
CELLS.update({f'AA{i}': v for i, v in enumerate([250.4, 249.1, 251.8, 250.9, 248.7, 252.2, 250.0, 249.6, 251.1, 250.7], start=1)})
CELLS.update({f'AB{i}': v for i, v in enumerate([248.9, 251.6, 247.8, 252.4, 249.0, 250.2, 246.9, 251.9, 249.4, 250.8, 248.2, 252.9], start=1)})
CELLS.update({'AD1': 48, 'AE1': 31, 'AF1': 21, 'AD2': 62, 'AE2': 29, 'AF2': 40})
CELLS.update({'AP1': 100, 'AP2': 131, 'AQ1': 110, 'AR1': 60, 'AS1': 61, 'AT1': 231})  # AD1:AF2's row totals, column totals, total
CELLS.update({'AD4': 100 * 110 / 231, 'AE4': 100 * 60 / 231, 'AF4': 100 * 61 / 231, 'AD5': 131 * 110 / 231, 'AE5': 131 * 60 / 231, 'AF5': 131 * 61 / 231})
CELLS.update({'AH1': 38, 'AI1': 41, 'AJ1': 29, 'AK1': 33, 'AL1': 27, 'AM1': 36, 'AN1': 27})
CELLS.update({'AH2': 33, 'AI2': 33, 'AJ2': 33, 'AK2': 33, 'AL2': 33, 'AM2': 33, 'AN2': 33})

FORMULAS = [
    # operators and precedence
    '=-2^2', '=2^3^2', '=10-2-3', '=2*3%', '=50%*20', '=-3^-2', '=(1+2)*3', '=1+2*3-4/2', '="3"+4', '="abc"+1',
    '=TRUE+TRUE', '=1/0', '=0.1+0.2=0.3', '=1=1', '="a"="A"', '="a"<"b"', '=1<"a"', '=TRUE>"z"', '=H4=0', '=H4=""',
    '="a"&TRUE', '=1&2', '="Total: "&F2', '=H2*2', '=H5*1', '=-H2', '="5"*"2"', '=H3*3', '=2.5&""',
    # aggregates and the rule that references skip text and logicals
    '=SUM(F2:F13)', '=SUM(F:F)', '=SUM(H1:H8)', '=SUM(H2)', '=SUM("5",TRUE)', '=SUM(H1:H8)+H2', '=AVERAGE(H1:H8)',
    '=COUNT(H1:H8)', '=COUNTA(H1:H8)', '=COUNTBLANK(H1:H8)', '=MAX(F2:F13)', '=MIN(F2:F13)', '=MEDIAN(F2:F13)',
    '=AVERAGE(F2:F13)', '=STDEV.S(F2:F13)', '=STDEV.P(F2:F13)', '=VAR.S(F2:F13)', '=VAR.P(F2:F13)', '=STDEV(D2:D13)',
    '=PERCENTILE.INC(F2:F13,0.9)', '=QUARTILE.INC(F2:F13,1)', '=QUARTILE.INC(F2:F13,3)', '=LARGE(F2:F13,2)',
    '=SMALL(F2:F13,3)', '=RANK.EQ(F5,F2:F13)', '=RANK.EQ(F5,F2:F13,1)', '=MODE.SNGL(S1:S8)', '=PRODUCT(D2:D4)',
    '=AVERAGEA(H1:H8)', '=MAX(H1:H8)', '=MIN(H1:H8)', '=AVERAGE(H4)',
    '=CORREL(D2:D13,F2:F13)', '=RSQ(F2:F13,D2:D13)', '=SLOPE(F2:F13,D2:D13)', '=INTERCEPT(F2:F13,D2:D13)',
    '=FORECAST.LINEAR(12,F2:F13,D2:D13)', '=STEYX(F2:F13,D2:D13)', '=COVARIANCE.S(D2:D13,F2:F13)',
    '=COVARIANCE.P(D2:D13,F2:F13)', '=INDEX(LINEST(F2:F13,D2:D13),1)', '=INDEX(LINEST(F2:F13,D2:D13),2)',
    '=INDEX(LINEST(F2:F13,D2:D13,TRUE,TRUE),2,1)', '=INDEX(LINEST(F2:F13,D2:D13,TRUE,TRUE),3,1)',
    '=INDEX(LINEST(F2:F13,D2:D13,TRUE,TRUE),4,1)', '=INDEX(LINEST(F2:F13,D2:D13,TRUE,TRUE),5,2)',
    # conditional sums and counts, criteria and wildcards
    '=SUMIF(B2:B13,"Online",F2:F13)', '=SUMIF(B2:B13,"<>Online",F2:F13)', '=SUMIF(D2:D13,">=3",F2:F13)',
    '=SUMIF(C2:C13,"*Blend",F2:F13)', '=SUMIF(C2:C13,"?ouse*",F2:F13)', '=SUMIF(F2:F13,">20")',
    '=COUNTIF(B2:B13,"online")', '=COUNTIF(D2:D13,3)', '=COUNTIF(D2:D13,"3")', '=COUNTIF(D2:D13,"<"&D3)',
    '=COUNTIF(H1:H8,"")', '=COUNTIF(H1:H8,"<>")', '=COUNTIF(H1:H8,">5")', '=COUNTIF(H1:H8,TRUE)',
    '=COUNTIF(H1:H8,"*")', '=COUNTIF(A2:A13,">="&DATE(2026,3,15))', '=COUNTIF(C2:C13,"*a*")',
    '=SUMIFS(F2:F13,B2:B13,"Online",C2:C13,"House Blend")', '=SUMIFS(F2:F13,A2:A13,">="&DATE(2026,3,10),A2:A13,"<"&DATE(2026,3,20))',
    '=COUNTIFS(B2:B13,"Café",D2:D13,">2")', '=AVERAGEIF(B2:B13,"Wholesale",F2:F13)', '=AVERAGEIFS(F2:F13,B2:B13,"Online",D2:D13,">1")',
    '=MAXIFS(F2:F13,B2:B13,"Café")', '=MINIFS(F2:F13,B2:B13,"Online")', '=AVERAGEIF(B2:B13,"Retail",F2:F13)',
    '=SUMPRODUCT(D2:D13,E2:E13)', '=SUMPRODUCT((B2:B13="Online")*F2:F13)', '=SUMPRODUCT(--(B2:B13="Café"))',
    '=SUMIF(B2:B13,"Café",D2:D13)/COUNTIF(B2:B13,"Café")',
    # logic
    '=IF(F2>20,"Large","Small")', '=IF(F5>100,"High",IF(F5>20,"Mid","Low"))', '=IF(H4,"yes","no")', '=IF(FALSE,1)',
    '=IFS(D2>10,"bulk",D2>1,"multi",TRUE,"single")', '=SWITCH(B3,"Online","web","Café","shop","other")',
    '=SWITCH(B5,"Online","web","Café","shop")', '=AND(D2>1,E2>10)', '=OR(D3>5,E3>5)', '=NOT(D2>1)', '=AND(H1:H8)',
    '=IFERROR(1/0,"none")', '=IFERROR(F2/D2,0)', '=IFNA(VLOOKUP("Tea",C2:F13,4,FALSE),0)', '=ISBLANK(H4)', '=ISBLANK(H8)',
    '=ISNUMBER(H2)', '=ISTEXT(H2)', '=ISERROR(H5*1)', '=ISNA(NA())', '=IF(H2=5,"same","different")', '=H2="5"',
    # maths
    '=ROUND(2.345,2)', '=ROUND(-1.5,0)', '=ROUND(1234.5,-2)', '=ROUND(1.005,2)', '=ROUND(F3,0)', '=ROUNDUP(1.21,1)',
    '=ROUNDDOWN(-1.29,1)', '=ROUNDUP(-1.21,1)', '=INT(-2.5)', '=TRUNC(-2.5)', '=TRUNC(3.789,2)', '=MOD(-7,3)', '=MOD(7,-3)',
    '=MOD(10,4)', '=ABS(-4)', '=SQRT(-1)', '=SQRT(16)', '=POWER(2,10)', '=CEILING(4.2,0.5)', '=FLOOR(4.7,0.5)',
    '=ROUND(F2/3,2)*3', '=SIGN(-0.2)', '=EXP(1)', '=LN(10)', '=LOG10(1000)',
    # finance (the F04 functions)
    '=PMT(0.05/12,60,20000)', '=PMT(0,12,1200)', '=FV(0.04,10,-1000)', '=FV(0.04/12,24,-100,-500,1)', '=PV(0.06,5,-2000)',
    '=NPV(0.08,-1000,300,400,500)', '=PMT(0.06,4,-20000,0)',
    # text
    '=LEN(C2)', '=UPPER(C2)', '=LOWER(B3)', '=PROPER("lantern COFFEE co.")', '=TRIM(H7)', '=LEN(TRIM(H7))',
    '=LEFT(C2,5)', '=RIGHT(C2,5)', '=MID(C2,7,3)', '=LEFT(C2)', '=CONCAT(B2," – ",C2)', '=CONCATENATE(B2,"/",D2)',
    '=TEXTJOIN(", ",TRUE,J1,K1,H4,M1)', '=TEXTJOIN("-",FALSE,H3,H4,H8)', '=SUBSTITUTE(C2,"Blend","Roast")',
    '=SUBSTITUTE("a-b-c","-","+",2)', '=FIND("B",C2)', '=FIND("b",C2)', '=SEARCH("b",C2)', '=SEARCH("?lend",C2)',
    '=EXACT("Online",B2)', '=EXACT("online",B2)', '=VALUE("12.5")', '=VALUE(H2)+1', '=VALUE("abc")',
    '=TEXT(1234.567,"#,##0.00")', '=TEXT(0.256,"0.0%")', '=TEXT(A2,"mmm yyyy")', '=TEXT(A2,"dddd")',
    '=TEXT(A2,"d mmmm yyyy")', '=TEXT(-1234.5,"#,##0")', '=TEXT(7,"000")', '=REPT("*",3)', '=LEFT(F2,1)', '=LEN(F3)',
    # dates
    '=DATE(2026,3,14)', '=DATE(2026,13,1)', '=DATE(2026,3,0)', '=YEAR(A2)', '=MONTH(A13)', '=DAY(A13)',
    '=EOMONTH(A2,0)', '=EOMONTH(A2,-1)', '=EDATE(DATE(2026,1,31),1)', '=WEEKDAY(A2)', '=WEEKDAY(A2,2)',
    '=DAYS(A13,A2)', '=A13-A2', '=DATEVALUE("2026-03-14")', '=DATE(2024,2,29)+365', '=MONTH(EOMONTH(A2,11))',
    '=DATE(1900,3,1)', '=DATE(1900,1,1)', '=YEAR(60)', '=DAY(60)',
    # lookup and reference
    '=VLOOKUP("Pastry",C2:F13,4,FALSE)', '=VLOOKUP("Tea",C2:F13,4,FALSE)', '=VLOOKUP("Pastry",C2:F13,5,FALSE)',
    '=VLOOKUP(180,J2:K5,2,TRUE)', '=VLOOKUP(500,J2:K5,2)', '=VLOOKUP(-5,J2:K5,2,TRUE)', '=VLOOKUP(99.99,J2:K5,2,TRUE)',
    '=VLOOKUP("Dec*",C2:F13,3,FALSE)', '=VLOOKUP("house blend",C2:F13,2,FALSE)', '=HLOOKUP("Q3",M1:P2,2,FALSE)',
    '=MATCH("Pastry",C2:C13,0)', '=MATCH(260,J2:J5,1)', '=MATCH(60,R1:R4,-1)', '=MATCH("Tea",C2:C13,0)',
    '=INDEX(F2:F13,MATCH("Single Origin",C2:C13,0))', '=INDEX(A1:F13,4,3)', '=INDEX(M2:P2,3)', '=INDEX(J2:K5,3,2)',
    '=CHOOSE(2,"a","b","c")', '=ROW(C5)', '=COLUMN(C5)', '=ROWS(A2:F13)', '=COLUMNS(A2:F13)',
    '=SUM(INDEX(F2:F13,1):INDEX(F2:F13,3))', '=INDEX(F2:F13,MATCH(MAX(F2:F13),F2:F13,0))',
    '=VLOOKUP(INDEX(C2:C13,5),C2:D13,2,FALSE)',
    # text that reads as a number or a date (US settings), and more date functions
    '=VALUE("$58.00")', '="$1,234.50"*1', '=VALUE("(5.00)")', '=VALUE("12%")', '=VALUE("1,23")', '=DATEVALUE("March 2, 2026")',
    '=DATEVALUE("22-MAY-2011")', '=DATEVALUE("2011/02/23")', '=DATEVALUE("3/14/2026")', '=DATEVALUE("2026-02-30")', '="2026-03-14"+1',
    '=NETWORKDAYS(A2,A13)', '=NETWORKDAYS(A2,A13,A7)', '=WORKDAY(A2,10)', '=WORKDAY(A13,-3)', '=DATEDIF(A2,DATE(2027,3,1),"Y")',
    '=DATEDIF(A2,DATE(2027,3,1),"M")', '=DATEDIF(A2,A13,"D")', '=WEEKNUM(A2)', '=WEEKNUM(A2,2)', '=WEEKNUM(DATE(2026,1,4))',
    '=LEN(TRIM("  a  b "&CHAR(160)))', '=PROPER("o\'neil 2-way STREET")', '=CODE("A")', '=CHAR(65)', '=CODE(CHAR(160))',
    '=SUMIF(B2:B13,1/0,F2:F13)', '=COUNTIF(B2:B13,NA())', '=SUMIFS(F2:F13,B2:B13,1/0)',
    '=PERCENTILE.EXC(F2:F13,0.25)', '=QUARTILE.EXC(F2:F13,1)', '=QUARTILE.EXC(F2:F13,3)', '=PERCENTILE.EXC(F2:F13,0.05)',
    '=SKEW(F2:F13)', '=SKEW.P(F2:F13)', '=KURT(F2:F13)', '=SKEW(D2:D13)', '=KURT(S1:S8)', '=KURT(1,2,3)',
    '=1.333+1.225-1.333-1.225', '=(1.333+1.225)-(1.333+1.225)', '=0.1+0.2-0.3', '=10.1-10', '=COUNTIF(H1:H8,H4)',
    '=SUMIF(B2:B13,"Online",F2:F5)', '=AVERAGEIF(B2:B13,"Café",F2)', '=COUNTIFS(D2:D13,H4)',
    '=CEILING.MATH(4.2)', '=CEILING.MATH(4.2,0.5)', '=CEILING.MATH(-4.2)', '=CEILING.MATH(-4.2,1,1)', '=CEILING.MATH(45/40)',
    '=FLOOR.MATH(4.7)', '=FLOOR.MATH(-4.2)', '=FLOOR.MATH(-4.2,1,1)', '=FLOOR.MATH(17,5)',
    '=INDEX(TREND(F2:F13,D2:D13,{5;12}),2)', '=INDEX(TREND(F2:F13,D2:D13,{5;12},FALSE),1)',
    '=INDEX(LINEST(F2:F13,D2:D13,FALSE),1)', '=INDEX(LINEST(F2:F13,D2:D13,FALSE),2)', '=INDEX(LINEST(F2:F13,D2:D13,FALSE,TRUE),2,1)',
    '=INDEX(LINEST(F2:F13,D2:D13,FALSE,TRUE),2,2)', '=INDEX(LINEST(F2:F13,D2:D13,FALSE,TRUE),3,1)', '=INDEX(LINEST(F2:F13,D2:D13,FALSE,TRUE),3,2)',
    '=INDEX(LINEST(F2:F13,D2:D13,FALSE,TRUE),4,1)', '=INDEX(LINEST(F2:F13,D2:D13,FALSE,TRUE),4,2)', '=INDEX(LINEST(F2:F13,D2:D13,FALSE,TRUE),5,1)',
    '=INDEX(LINEST(F2:F13,D2:D13,FALSE,TRUE),5,2)', '=INDEX(LINEST(F2:F13),1)', '=INDEX(LINEST(F2:F13,,TRUE,TRUE),3,1)',
    '=WEEKDAY(A2,11)', '=WEEKDAY(A2,12)', '=WEEKDAY(A2,13)', '=WEEKDAY(A3,14)', '=WEEKDAY(A3,15)', '=WEEKDAY(A3,16)', '=WEEKDAY(A3,17)',
    '=WEEKDAY(A2,4)', '=WEEKDAY(A2,0)', '=WEEKDAY(A2,3)',
    '=10^308*10', '=FORECAST.LINEAR(5,F2:F4,{2;2;2})', '=IFERROR(10^308*10,"big")', '=NUMBERVALUE("2,5",",")',
    '=NORM.DIST(1000,1004,3,TRUE)', '=NORM.DIST(1000,1004,3,FALSE)', '=NORM.DIST(1,0,0,TRUE)', '=NORM.S.DIST(-1.96,TRUE)', '=NORM.S.DIST(0.5,FALSE)',
    '=NORM.S.DIST(-8.2,TRUE)', '=NORM.INV(0.01,1004,3)', '=NORM.INV(1,0,1)', '=NORM.S.INV(0.975)', '=STANDARDIZE(12.6,39.38,26.1)', '=STANDARDIZE(1,2,0)',
    '=BINOM.DIST(3,10,0.3,FALSE)', '=BINOM.DIST(3,10,0.3,TRUE)', '=BINOM.DIST(3.9,10.2,0.3,TRUE)', '=BINOM.DIST(11,10,0.3,TRUE)', '=BINOM.INV(200,0.31,0.025)', '=BINOM.INV(12,0.5,0.99)', '=POISSON.DIST(2,1.9,FALSE)', '=POISSON.DIST(2.7,1.9,TRUE)',
    '=POISSON.DIST(-1,1.9,TRUE)', '=EXPON.DIST(0.5,2,TRUE)', '=EXPON.DIST(3,0.4,FALSE)', '=T.DIST(60,1,TRUE)', '=T.DIST(8,3,FALSE)', '=T.DIST(-2.1,9,TRUE)',
    '=T.DIST(2.1,0.5,TRUE)', '=T.DIST.RT(2.1,9)', '=T.DIST.RT(-1.3,4)', '=T.DIST.2T(12,4)', '=T.TEST(AA1:AA10,AB1:AB12,2,3)', '=T.TEST(AA1:AA10,AB1:AB12,1,2)',
    '=T.TEST(AA1:AA10,AB1:AB10,2,1)', '=T.TEST(AA1:AA10,AB1:AB12,1,1)', '=T.TEST(AA1:AA10,AB1:AB12,3,2)', '=T.TEST(AA1:AA10,AB1:AB12,2.7,3.2)',
    '=Z.TEST(AA1:AA10,250)', '=Z.TEST(AA1:AA10,250,1.5)', '=CONFIDENCE.NORM(0.05,2.5,50)', '=CONFIDENCE.NORM(0.05,2.5,50.9)', '=CONFIDENCE(0.1,3,40)',
    '=CHISQ.DIST(3.84,1,TRUE)', '=CHISQ.DIST(3,4,FALSE)', '=CHISQ.DIST.RT(11.07,5)', '=CHISQ.DIST.RT(40,6.5)', '=CHISQ.INV(0.95,3)', '=CHISQ.INV.RT(0.05,6)',
    '=CHISQ.TEST(AD1:AF2,AD4:AF5)', '=CHISQ.TEST(AH1:AN1,AH2:AN2)', '=CHISQ.TEST(AD1:AF2,AH1:AN1)', '=F.DIST(3.1,3,41,TRUE)', '=F.DIST(1,5,5,FALSE)',
    '=F.DIST.RT(3.1,3,41)', '=F.INV(0.05,2,10)', '=F.INV.RT(0.05,3,41)', '=F.TEST(AA1:AA10,AB1:AB12)', '=COMBIN(52,5)', '=COMBIN(8.9,2.2)', '=COMBIN(2,3)',
    '=PERMUT(100,3)', '=FACT(10)', '=FACT(5.9)', '=FACT(-1)', '=DEVSQ(AA1:AA10)', '=SUMSQ(3,4)', '=AVEDEV(AA1:AA10)', '=GEOMEAN(AA1:AA10)', '=HARMEAN(AA1:AA10)',
    '=TRIMMEAN({4,5,6,7,2,3,4,5,1,2,3},0.2)', '=TRIMMEAN(AB1:AB12,0.25)', '=INDEX(MODE.MULT(S1:S8),1)', '=INDEX(MODE.MULT(S1:S8),2)', '=ROWS(MODE.MULT(S1:S8))',
    '=PERCENTRANK.INC(S1:S8,5)', '=PERCENTRANK.INC(S1:S8,4.5)', '=PERCENTRANK.INC(S1:S8,4.5,2)', '=PERCENTRANK.EXC(S1:S8,5)', '=PERCENTRANK.EXC(S1:S8,6.2)',
    '=PERCENTRANK.INC(AA1:AA10,250.5)', '=PERCENTRANK.INC(S1:S8,3)', '=PERCENTRANK.EXC(S1:S8,2.5,4)', '=SUM(LARGE(S1:S8,{1,2}))', '=SUM(SMALL(AA1:AA10,{1;2;3}))', '=SUM(BINOM.DIST({0,1},5,0.5,FALSE))',
    '=SUM(POISSON.DIST({0,1,2},2,FALSE))', '=INDEX(NORM.S.DIST({-1,0,1},TRUE),1,3)', '=SUMPRODUCT(T.DIST({-1,1},5,TRUE))',
    '=CHISQ.TEST(AD1:AF2,MMULT(AP1:AP2,AQ1:AS1)/AT1)', '=SUM(MMULT({1,2,3},{4;5;6}))', '=INDEX(MMULT({1,2;3,4},{5;6}),2,1)',
    '=INDEX(MMULT({1;2},{3,4}),2,2)', '=ROWS(MMULT(AP1:AP2,AQ1:AS1))', '=COLUMNS(MMULT(AP1:AP2,AQ1:AS1))', '=MMULT({1,2},{1,2})',
    '=SUM(MMULT(AP1:AP3,{1}))', '=SUM(MMULT({1,"a"},{1;2}))',
    '=SUM(ROW(12:20))', '=ROWS(12:20)', '=SUM(BINOM.DIST(ROW(12:20),20,0.45,FALSE))', '=SUM(ROW($3:$5))',
    '=VLOOKUP(13,J1:K40,2,TRUE)', '=VLOOKUP(260,J:K,2,TRUE)', '=MATCH(120,J1:J40,1)', '=VLOOKUP(-5,J1:K40,2,TRUE)', '=T.INV.2T(0.05,10)', '=T.INV(0.975,30)',
    '=T.DIST.2T(2.1,15)', '=CONFIDENCE.T(0.05,STDEV.S(F2:F13),COUNT(F2:F13))', '=CONFIDENCE.T(0.05,1,1)',
    '=SUM(D2:F2 E1:E5)', '=SUM(A1:A2 C1:C2)', '=N(TRUE)+N("7")+N(5)', '=NUMBERVALUE("2.500,27",",",".")', '=NUMBERVALUE("3.5%")',
    '=MROUND(17,5)', '=MROUND(F3,0.05)', '=LOG(8,2)', '=LOG(1000)', '=ISERR(NA())', '=ISERR(1/0)', '=ISLOGICAL(H3)',
]
# Functions written with Excel's "_xlfn." prefix in the file format (introduced after Excel 2007)
XLFN = ['NUMBERVALUE', 'T.INV.2T', 'T.INV', 'T.DIST.2T', 'CONFIDENCE.T', 'CEILING.MATH', 'FLOOR.MATH', 'PERCENTILE.EXC', 'QUARTILE.EXC', 'SKEW.P', 'STDEV.S', 'STDEV.P', 'VAR.S', 'VAR.P', 'PERCENTILE.INC', 'QUARTILE.INC', 'RANK.EQ', 'MODE.SNGL', 'FORECAST.LINEAR',
        'COVARIANCE.S', 'COVARIANCE.P', 'MAXIFS', 'MINIFS', 'IFS', 'SWITCH', 'IFNA', 'CONCAT', 'TEXTJOIN', 'DAYS',
        'NORM.DIST', 'NORM.S.DIST', 'NORM.INV', 'NORM.S.INV', 'BINOM.DIST.RANGE', 'BINOM.DIST', 'BINOM.INV', 'POISSON.DIST', 'EXPON.DIST', 'T.DIST.RT', 'T.DIST',
        'T.TEST', 'Z.TEST', 'CONFIDENCE.NORM', 'CHISQ.DIST.RT', 'CHISQ.DIST', 'CHISQ.INV.RT', 'CHISQ.INV', 'CHISQ.TEST', 'F.DIST.RT', 'F.DIST', 'F.INV.RT', 'F.INV', 'F.TEST', 'MODE.MULT', 'PERCENTRANK.INC', 'PERCENTRANK.EXC']
# Cases where LibreOffice gives a different answer from Excel: the test uses Excel's value, with the reason.
# (a) LibreOffice has no logical type: TRUE is the number 1, so it sums, counts and joins as 1. Excel's COUNT/SUM
#     help: "Logical values ... typed directly into the list of arguments are counted"; in a reference "only numbers
#     ... are counted. Empty cells, logical values, text, or error values in the array or reference are not counted";
#     =SUM("5", 15, TRUE) is 21. Excel orders numbers < text < FALSE < TRUE when comparing, and TRUE joins as "TRUE".
# (b) Excel's COUNTIF help: =COUNTIF(A2:A5,"*") "Counts the number of cells containing any text".
# (c) Excel's SQRT help: "If number is negative, SQRT returns the #NUM! error value."
# (d) Excel's VLOOKUP help: "If col_index_num is greater than the number of columns in table-array, you'll get the
#     #REF! error value."
# (e) Excel treats 1900 as a leap year (Microsoft: "Excel incorrectly assumes that the year 1900 is a leap year"), so
#     1 January 1900 is serial 1 and serial 60 is 29 February 1900; LibreOffice counts from 30 December 1899.
# (f) Excel's DATEVALUE help: =DATEVALUE("2011/02/23") is 40597 (LibreOffice in English (USA) doesn't read year/month/day
#     with slashes).
# (g) Excel's CODE help: the code "corresponds to the character set used by your computer" (Windows-1252 on Windows,
#     where CHAR(160) is the non-breaking space and its code is 160); LibreOffice reports 239 here.
# (h) Excel's PERCENTILE.EXC help: "If it cannot interpolate for the percentile, k specified, Excel will return #NUM!
#     error" (here k = 0.05 with 12 values, below 1/(n + 1)); LibreOffice returns #VALUE!.
# (i) Excel's NUMBERVALUE help: =NUMBERVALUE("3.5%") returns 0.035 ("If the Text argument ends in one or more percent
#     signs (%), they are used in the calculation of the result"); LibreOffice returns #VALUE!.
# (j) Excel's WEEKDAY help: "If return_type is out of the range specified in the table above, a #NUM! error is
#     returned"; LibreOffice returns #VALUE! for return types 4 and 0.
# (k) Excel's help documents #NUM! for arguments outside a function's domain: NORM.DIST and STANDARDIZE "If
#     standard_dev ≤ 0", NORM.INV "If probability <= 0 or if probability >= 1", BINOM.DIST "If number_s < 0 or number_s >
#     trials", POISSON.DIST "If x < 0", COMBIN "number < number_chosen", FACT "If number is negative"; T.TEST "#NUM!" when
#     tails is not 1 or 2 and "#N/A" for a paired test (type 1) with different numbers of data points; CHISQ.TEST "#N/A"
#     when the ranges have different numbers of data points; T.DIST "an error value" when deg_freedom < 1 (the engine
#     returns #NUM!, as the other t functions do). LibreOffice returns #VALUE! (or #DIV/0! for STANDARDIZE) in these cases.
# (l) Excel's PERCENTRANK.INC/.EXC cut the result to the significant digits without rounding: Microsoft's examples give
#     0.555 for 5/9 and 0.666 for 6/9; LibreOffice rounds (0.556, 0.536, 0.54, 0.286, 0.2778).
# (m) Excel 365 evaluates a function once per element when an argument that takes one value is given an array, and the
#     results spill (Microsoft, "Guidelines and examples of array formulas": =SMALL(B9:B18,{1;2;3}) evaluates SMALL three
#     times); LibreOffice, outside an array formula, uses only the first element. The expected value is SciPy's
#     binom.sf(11, 20, 0.45), the sum over 12 to 20 successes.
EXCEL_DIFFERS = {
    '=SUM(BINOM.DIST(ROW(12:20),20,0.45,FALSE))': __import__('scipy.stats', fromlist=['binom']).binom.sf(11, 20, 0.45).item(),   # (m)
    '=PERCENTRANK.INC(S1:S8,4.5)': 0.535, '=PERCENTRANK.INC(S1:S8,4.5,2)': 0.53, '=PERCENTRANK.EXC(S1:S8,5)': 0.555,       # (l)
    '=PERCENTRANK.INC(S1:S8,3)': 0.285, '=PERCENTRANK.EXC(S1:S8,2.5,4)': 0.2777,
    '=NORM.DIST(1,0,0,TRUE)': {'error': '#NUM!'}, '=NORM.INV(1,0,1)': {'error': '#NUM!'}, '=STANDARDIZE(1,2,0)': {'error': '#NUM!'},   # (k)
    '=BINOM.DIST(11,10,0.3,TRUE)': {'error': '#NUM!'}, '=POISSON.DIST(-1,1.9,TRUE)': {'error': '#NUM!'}, '=T.DIST(2.1,0.5,TRUE)': {'error': '#NUM!'},
    '=T.TEST(AA1:AA10,AB1:AB12,1,1)': {'error': '#N/A'}, '=T.TEST(AA1:AA10,AB1:AB12,3,2)': {'error': '#NUM!'}, '=CHISQ.TEST(AD1:AF2,AH1:AN1)': {'error': '#N/A'},
    '=COMBIN(2,3)': {'error': '#NUM!'}, '=FACT(-1)': {'error': '#NUM!'},
    '=WEEKDAY(A2,4)': {'error': '#NUM!'}, '=WEEKDAY(A2,0)': {'error': '#NUM!'},                           # (j)
    '=NUMBERVALUE("3.5%")': 0.035,                                                                         # (i)
    '=DATEVALUE("2011/02/23")': 40597,                                                                     # (f)
    '=PERCENTILE.EXC(F2:F13,0.05)': {'error': '#NUM!'},                                                     # (h)
    '=CODE(CHAR(160))': 160,                                                                               # (g)
    '=TRUE>"z"': True, '="a"&TRUE': 'aTRUE', '=SUM(H1:H8)': 6.5, '=SUM("5",TRUE)': 6, '=SUM(H1:H8)+H2': 11.5,
    '=AVERAGE(H1:H8)': 6.5 / 3, '=COUNT(H1:H8)': 3, '=TEXTJOIN("-",FALSE,H3,H4,H8)': 'TRUE--0',      # (a)
    '=COUNTIF(H1:H8,"*")': 3,                                                                            # (b)
    '=SQRT(-1)': {'error': '#NUM!'},                                                                     # (c)
    '=VLOOKUP("Pastry",C2:F13,5,FALSE)': {'error': '#REF!'},                                             # (d)
    '=DATE(1900,1,1)': 1, '=DAY(60)': 29,                                                                # (e)
}

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
  Sub RecalculateAndSave()
    ThisComponent.calculateAll()
    ThisComponent.store()
    ThisComponent.close(True)
  End Sub
</script:module>"""


def xlfn(f):
    import re
    for name in XLFN:
        f = re.sub(r'(?<![A-Za-z_.])' + re.escape(name) + r'\(', '_xlfn.' + name + '(', f)
    return f


def recalc(path):
    if len(sys.argv) > 1:
        subprocess.run([sys.executable, sys.argv[1], str(path), '60'], check=True, capture_output=True)
        return
    prof = Path(tempfile.mkdtemp())
    url = prof.as_uri()
    subprocess.run(['soffice', '--headless', '--terminate_after_init', f'-env:UserInstallation={url}'], capture_output=True, timeout=120)
    mdir = prof / 'user' / 'basic' / 'Standard'
    (mdir / 'Module1.xba').write_text(MACRO)
    subprocess.run(['soffice', '--headless', f'-env:UserInstallation={url}',
                    'vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application',
                    str(path)], capture_output=True, timeout=180, check=True)
    shutil.rmtree(prof, ignore_errors=True)


def main():
    tmp = Path(tempfile.mkdtemp()) / 'sheet.xlsx'
    wb = Workbook(); ws = wb.active
    for k, v in CELLS.items(): ws[k] = v
    at = {}
    for i, f in enumerate(FORMULAS, start=1):
        a = f'Z{i}'; at[f] = a
        ws[a] = xlfn(f)
    wb.save(tmp)
    recalc(tmp)
    vals = load_workbook(tmp, data_only=True).active
    out = []
    for f in FORMULAS:
        v = vals[at[f]].value
        if isinstance(v, str) and v.startswith('#'): v = {'error': v}
        elif isinstance(v, (dt.datetime, dt.date)): v = (v.date() if isinstance(v, dt.datetime) else v) - dt.date(1899, 12, 30); v = v.days
        out.append([at[f], f, EXCEL_DIFFERS[f], {'libreoffice': v}] if f in EXCEL_DIFFERS else [at[f], f, v])
    fx = {'source': 'LibreOffice ' + subprocess.run(['soffice', '--version'], capture_output=True, text=True).stdout.split()[1],
          'cells': CELLS, 'cases': out}
    (HERE / 'fixtures-sheet.json').write_text(json.dumps(fx, ensure_ascii=False, indent=0))
    print(f'{len(out)} cases → tests/fixtures-sheet.json')


if __name__ == '__main__':
    main()
