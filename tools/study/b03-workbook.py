#!/usr/bin/env python3
"""Build assets/data/study/lantern-shop-workbook.xlsx, the B03 practice workbook, from assets/js/data/study/b03/case.js.

Run:  python3 tools/study/b03-workbook.py
Needs node, openpyxl and LibreOffice (soffice), which recalculates the formulas so the file opens with their values.
The data are the course's: Lantern's 231 cleaned spring orders, the price list, the delivery and size bands, the 76
catering events, the July stock plan (laid out for Solver) and the monthly shop model (for Goal Seek and data tables).
"""
import json, shutil, subprocess, sys, tempfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import datetime as dt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'assets/data/study/lantern-shop-workbook.xlsx'
JS = """
const c = await import('%s');
console.log(JSON.stringify({ ORDER_COLUMNS: c.ORDER_COLUMNS, ORDERS: c.ORDERS, EVENT_COLUMNS: c.EVENT_COLUMNS, EVENTS: c.EVENTS,
  PRODUCT_COLUMNS: c.PRODUCT_COLUMNS, PRODUCTS: c.PRODUCTS, DELIVERY_BANDS: c.DELIVERY_BANDS, SIZE_BANDS: c.SIZE_BANDS,
  STOCK_PLAN: c.STOCK_PLAN, SHOP_SHEET: c.SHOP_SHEET, QUOTE: c.QUOTE }));
""" % (ROOT / 'assets/js/data/study/b03/case.js').as_uri()
D = json.loads(subprocess.run(['node', '--input-type=module', '-e', JS], capture_output=True, text=True, check=True).stdout)

HEAD = Font(bold=True)
FILL = PatternFill('solid', fgColor='F3ECDB')
INPUT = PatternFill('solid', fgColor='FBF0D2')
MONEY, DATE = '$#,##0.00', 'd mmm yyyy'

def table(ws, name, columns, rows, formats={}):
    ws.append(columns)
    for r in rows:
        ws.append([dt.date.fromisoformat(v) if isinstance(v, str) and len(v) == 10 and v[4] == '-' and v[7] == '-' else v for v in r])
    for j, c in enumerate(columns, start=1):
        ws.cell(1, j).font = HEAD
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(28, len(c) + 4, max(len(str(r[j - 1] or '')) for r in rows) + 3))
        fmt = formats.get(c)
        if fmt:
            for i in range(2, len(rows) + 2): ws.cell(i, j).number_format = fmt
    ref = f'A1:{get_column_letter(len(columns))}{len(rows) + 1}'
    t = Table(displayName=name, ref=ref); t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = 'A2'

wb = Workbook()
about = wb.active; about.title = 'About'
for line in [
    ['Lantern Coffee Co. — B03 Excel for Business Analytics practice workbook'],
    ['Every business, person and number in this workbook is fictional, invented for teaching.'],
    [],
    ['Sheet', 'What it holds'],
    ['Orders', 'The online shop’s 231 orders for 1 March – 30 June 2026 after cleaning (B02). Table name: Orders.'],
    ['Products', 'Kilnworth Roasting’s spring price list to Lantern (unit_cost) and Lantern’s shop prices. Table name: Products.'],
    ['Bands', 'The July delivery rule (free over $35, otherwise $4.95) and Ana’s order-size bands, for approximate-match lookups.'],
    ['Events', 'Lantern’s 76 catering events in 2026. Table name: Events.'],
    ['Stock plan', 'The first July stock order from Kilnworth, laid out for Solver (decision cells shaded).'],
    ['Shop model', 'One month of the relaunched shop, for Goal Seek and data tables (inputs shaded).'],
]: about.append(line)
about['A1'].font = Font(bold=True, size=13); about['A4'].font = HEAD; about['B4'].font = HEAD
about.column_dimensions['A'].width = 14; about.column_dimensions['B'].width = 110

table(wb.create_sheet('Orders'), 'Orders', D['ORDER_COLUMNS'], D['ORDERS'], {'order_date': DATE, 'unit_price': MONEY, 'discount': MONEY, 'order_total': MONEY})
table(wb.create_sheet('Products'), 'Products', D['PRODUCT_COLUMNS'], D['PRODUCTS'], {'price': MONEY, 'unit_cost': MONEY})
bands = wb.create_sheet('Bands')
bands.append(['Order from', 'Delivery', '', 'Size from', 'Band'])
for i in range(max(len(D['DELIVERY_BANDS']), len(D['SIZE_BANDS']))):
    a = D['DELIVERY_BANDS'][i] if i < len(D['DELIVERY_BANDS']) else [None, None]
    b = D['SIZE_BANDS'][i] if i < len(D['SIZE_BANDS']) else [None, None]
    bands.append([a[0], a[1], None, b[0], b[1]])
for c in 'ABDE': bands[f'{c}1'].font = HEAD
for r in range(2, 6): bands[f'A{r}'].number_format = MONEY; bands[f'B{r}'].number_format = MONEY; bands[f'D{r}'].number_format = MONEY
bands.column_dimensions['A'].width = 15; bands.column_dimensions['D'].width = 12; bands.column_dimensions['E'].width = 18
table(wb.create_sheet('Events'), 'Events', D['EVENT_COLUMNS'], D['EVENTS'], {'event_date': DATE, 'invoice': MONEY})

sp = wb.create_sheet('Stock plan')
SP = D['STOCK_PLAN']; prices = {p[0]: p for p in D['PRODUCTS']}
sp.append(['First July stock order (1–14 July 2026)'])
sp.append([])
sp.append(['Product', 'Bags to order', 'Margin per bag', 'Cost per bag', 'kg per bag', 'Packing minutes', 'At least', 'At most'])
for it in SP['items']:
    p = prices[it['product']]
    sp.append([it['product'], 0, round(p[3] - p[4], 2), p[4], it['kg'], it['minutes'], it['min'], it['max']])
n = len(SP['items']); last = 3 + n
sp.append([])
sp.append(['', 'Used', '', 'Limit'])
rows = [('Roasted coffee (kg)', f'=SUMPRODUCT(B4:B{last},E4:E{last})', SP['limits']['kg']),
        ('Stock budget ($)', f'=SUMPRODUCT(B4:B{last},D4:D{last})', SP['limits']['budget']),
        ('Packing time (minutes)', f'=SUMPRODUCT(B4:B{last},F4:F{last})', SP['limits']['minutes'])]
for label, f, lim in rows: sp.append([label, f, None, lim])
sp.append([])
sp.append(['Total contribution', f'=SUMPRODUCT(B4:B{last},C4:C{last})'])
sp['A1'].font = Font(bold=True, size=12)
for j in range(1, 9): sp.cell(3, j).font = HEAD
for i in range(4, last + 1): sp.cell(i, 2).fill = INPUT; sp.cell(i, 3).number_format = MONEY; sp.cell(i, 4).number_format = MONEY
sp.column_dimensions['A'].width = 28
for c in 'BCDEFGH': sp.column_dimensions[c].width = 15

sm = wb.create_sheet('Shop model')
for a, v in D['SHOP_SHEET']['cells'].items(): sm[a] = v
sm['A1'].font = Font(bold=True, size=12)
for r in range(3, 9): sm[f'B{r}'].fill = INPUT
for r in (4, 7, 8, 10, 11, 12, 13, 14, 15, 16): sm[f'B{r}'].number_format = MONEY
for r in (5, 6): sm[f'B{r}'].number_format = '0.0%'
sm.column_dimensions['A'].width = 36; sm.column_dimensions['B'].width = 14

tmp = Path(tempfile.mkdtemp()) / 'wb.xlsx'
wb.save(tmp)
# recalculate with LibreOffice so the formulas carry values when the file opens
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
  Sub RecalculateAndSave()
    ThisComponent.calculateAll()
    ThisComponent.store()
    ThisComponent.close(True)
  End Sub
</script:module>"""
prof = Path(tempfile.mkdtemp()); url = prof.as_uri()
subprocess.run(['soffice', '--headless', '--terminate_after_init', f'-env:UserInstallation={url}'], capture_output=True, timeout=120)
(prof / 'user/basic/Standard/Module1.xba').write_text(MACRO)
subprocess.run(['soffice', '--headless', f'-env:UserInstallation={url}', 'vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application', str(tmp)], capture_output=True, timeout=180, check=True)
shutil.rmtree(prof, ignore_errors=True)
shutil.copy(tmp, OUT)
print(f'wrote {OUT.relative_to(ROOT)} ({OUT.stat().st_size // 1024} KB)')
